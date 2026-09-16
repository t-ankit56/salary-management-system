from collections import Counter
from decimal import Decimal, InvalidOperation
from typing import BinaryIO

import openpyxl
from django.core.exceptions import ValidationError as DjangoValidationError
from django.db import transaction
from django.db.models import DateField

from employees.models import Country, Department, Employee, Role
from employees.services import create_employee

MAX_ROWS = 10_000

REQUIRED_COLUMNS = [
    "employee_code",
    "first_name",
    "last_name",
    "email",
    "department",
    "country",
    "role",
    "hire_date",
    "base",
    "allowance",
    "yearly_bonus",
    "salary_effective_from",
]


class RosterUploadError(Exception):
    def __init__(self, errors: list[dict]):
        self.errors = errors
        super().__init__("Roster upload validation failed")


def _parse_date(value):
    return DateField().to_python(value)


def _parse_amount(value):
    amount = Decimal(str(value))
    if amount < 0:
        raise ValueError("must not be negative")
    return amount


def _blank(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def import_roster(file: BinaryIO) -> dict:
    workbook = openpyxl.load_workbook(file, read_only=True, data_only=True)
    try:
        sheet = workbook.active

        header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        header_index = {name: idx for idx, name in enumerate(header)}
        missing_columns = [c for c in REQUIRED_COLUMNS if c not in header_index]
        if missing_columns:
            raise RosterUploadError(
                [{"row": 1, "error": f"missing required column: {c}"} for c in missing_columns]
            )

        total_data_rows = (sheet.max_row or 1) - 1
        if total_data_rows > MAX_ROWS:
            raise RosterUploadError(
                [{"row": None, "error": f"file exceeds the maximum of {MAX_ROWS} rows"}]
            )

        raw_rows = [
            (row_num, row)
            for row_num, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2)
            if any(cell is not None for cell in row)
        ]
    finally:
        workbook.close()

    errors: list[dict] = []
    parsed_rows = []
    codes_seen: Counter = Counter()
    emails_seen: Counter = Counter()

    for row_num, row in raw_rows:
        values = {col: row[idx] if idx < len(row) else None for col, idx in header_index.items()}
        values = {col: values.get(col) for col in REQUIRED_COLUMNS}

        for col, value in values.items():
            if _blank(value):
                errors.append({"row": row_num, "error": f"{col} is required"})

        if values["employee_code"]:
            codes_seen[values["employee_code"]] += 1
        if values["email"]:
            emails_seen[values["email"]] += 1

        parsed_rows.append((row_num, values))

    for row_num, values in parsed_rows:
        if values["employee_code"] and codes_seen[values["employee_code"]] > 1:
            errors.append({"row": row_num, "error": "duplicate employee_code in file"})
        if values["email"] and emails_seen[values["email"]] > 1:
            errors.append({"row": row_num, "error": "duplicate email in file"})

    departments = {d.name: d for d in Department.objects.all()}
    roles = {r.name: r for r in Role.objects.all()}
    countries = {c.name: c for c in Country.objects.all()}

    existing_by_code = {
        e.employee_code: e for e in Employee.objects.filter(employee_code__in=list(codes_seen))
    }
    existing_by_email = {e.email: e for e in Employee.objects.filter(email__in=list(emails_seen))}

    resolved_rows = []

    for row_num, values in parsed_rows:
        employee_code = values["employee_code"]
        email = values["email"]

        if email and email in existing_by_email:
            owner = existing_by_email[email]
            if owner.employee_code != employee_code:
                errors.append({"row": row_num, "error": "email belongs to a different employee"})

        department = departments.get(values["department"])
        if values["department"] and department is None:
            errors.append({"row": row_num, "error": "unknown department"})

        role = roles.get(values["role"])
        if values["role"] and role is None:
            errors.append({"row": row_num, "error": "unknown role"})

        country = countries.get(values["country"])
        if values["country"] and country is None:
            errors.append({"row": row_num, "error": "unknown country"})

        hire_date = None
        if values["hire_date"]:
            try:
                hire_date = _parse_date(values["hire_date"])
            except DjangoValidationError:
                errors.append({"row": row_num, "error": "hire_date does not parse"})

        salary_effective_from = None
        if values["salary_effective_from"]:
            try:
                salary_effective_from = _parse_date(values["salary_effective_from"])
            except DjangoValidationError:
                errors.append({"row": row_num, "error": "salary_effective_from does not parse"})

        amounts = {}
        for field in ("base", "allowance", "yearly_bonus"):
            if values[field] is None:
                continue
            try:
                amounts[field] = _parse_amount(values[field])
            except (InvalidOperation, ValueError):
                errors.append({"row": row_num, "error": f"{field} must be a non-negative number"})

        if hire_date and salary_effective_from and salary_effective_from < hire_date:
            errors.append(
                {"row": row_num, "error": "salary_effective_from must be on or after hire_date"}
            )

        resolved_rows.append(
            {
                "row_num": row_num,
                "employee_code": employee_code,
                "first_name": values["first_name"],
                "last_name": values["last_name"],
                "email": email,
                "department": department,
                "role": role,
                "country": country,
                "hire_date": hire_date,
                "salary_effective_from": salary_effective_from,
                "base": amounts.get("base"),
                "allowance": amounts.get("allowance"),
                "yearly_bonus": amounts.get("yearly_bonus"),
            }
        )

    if errors:
        raise RosterUploadError(errors)

    created = 0
    updated = 0

    with transaction.atomic():
        for row in resolved_rows:
            existing = existing_by_code.get(row["employee_code"])
            if existing is None:
                create_employee(
                    employee_code=row["employee_code"],
                    first_name=row["first_name"],
                    last_name=row["last_name"],
                    email=row["email"],
                    department=row["department"],
                    role=row["role"],
                    country=row["country"],
                    hire_date=row["hire_date"],
                    base=row["base"],
                    allowance=row["allowance"],
                    yearly_bonus=row["yearly_bonus"],
                    salary_effective_from=row["salary_effective_from"],
                )
                created += 1
            else:
                existing.first_name = row["first_name"]
                existing.last_name = row["last_name"]
                existing.email = row["email"]
                existing.department = row["department"]
                existing.role = row["role"]
                existing.country = row["country"]
                existing.hire_date = row["hire_date"]
                existing.save()
                updated += 1

    return {"created": created, "updated": updated}
